import argparse
import csv
import io
import re
import sys
import tempfile
import unicodedata
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt


REQUIRED_COLUMNS = ("通し番号", "金額", "班名")

IMAGE_EXTENSIONS = {
    ".jpg", ".jpeg", ".png",
    ".bmp", ".tif", ".tiff", ".webp",
}

# スライド設定：単位はインチ
SLIDE_WIDTH = 13.333333
SLIDE_HEIGHT = 7.5
MARGIN = 0.30
CONTENT_TOP = 0.90

HORIZONTAL_GAP = 0.10
VERTICAL_GAP = 0.12
CAPTION_HEIGHT = 0.22


def text(value):
    return "" if value is None else str(value).strip()


def fail_if_errors(errors):
    if errors:
        raise ValueError("\n".join(errors))


def parse_number(value, currency=False):
    """数値セル、または通貨文字列をDecimalへ変換する。"""
    if value is None or isinstance(value, bool):
        raise ValueError(f"数値が空欄または不正です: {value!r}")

    s = unicodedata.normalize("NFKC", str(value)).strip()

    if currency:
        s = s.replace("￥", "¥").replace("\\", "¥")

        pattern = (
            r"([+-]?)¥?([+-]?)"
            r"(\d{1,3}(?:,\d{3})+|\d+)"
            r"(\.\d+)?"
        )
        match = re.fullmatch(pattern, s)

        if not match or (match[1] and match[2]):
            raise ValueError(f"金額の形式が不正です: {value!r}")

        s = (match[1] or match[2]) + match[3].replace(",", "")
        s += match[4] or ""

    elif not re.fullmatch(r"\d+(?:\.0+)?", s):
        raise ValueError(f"番号の形式が不正です: {value!r}")

    result = Decimal(s)

    if not result.is_finite():
        raise ValueError(f"有限の数値ではありません: {value!r}")

    return result


def parse_receipt_id(value):
    number = parse_number(value)

    if number != number.to_integral_value() or number <= 0:
        raise ValueError(f"通し番号は正の整数が必要です: {value!r}")

    return int(number)


def read_excel(path, sheet_name, header_row, allow_cached_formulas):
    """
    通し番号ごとに金額最大の行を選び、その班名を採用する。

    最大金額が同額の場合：
      同じ班名 → Excelで先に出現した行を代表行にする
      異なる班名 → エラー
    """
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Excelは .xlsx または .xlsm を指定してください")

    formulas = load_workbook(
        path, data_only=False, read_only=True
    )
    cached = None

    try:
        cached = load_workbook(
            path, data_only=True, read_only=True
        )

        if sheet_name is None:
            if len(formulas.sheetnames) != 1:
                raise ValueError(
                    "複数シートがあります。--sheetで指定してください。\n"
                    f"シート一覧: {formulas.sheetnames}"
                )
            sheet_name = formulas.sheetnames[0]

        if sheet_name not in formulas.sheetnames:
            raise ValueError(f"シートがありません: {sheet_name}")

        ws = formulas[sheet_name]
        cached_ws = cached[sheet_name]

        header_cells = next(
            ws.iter_rows(
                min_row=header_row,
                max_row=header_row,
            ),
            (),
        )
        headers = [text(cell.value) for cell in header_cells]

        columns = {}

        for name in REQUIRED_COLUMNS:
            hits = [
                index
                for index, header in enumerate(headers)
                if header == name
            ]

            if len(hits) != 1:
                raise ValueError(
                    f"列「{name}」は1列だけ必要です。"
                    f"検出数: {len(hits)}"
                )

            columns[name] = hits[0]

        groups = defaultdict(list)
        errors = []
        formula_count = 0
        last_column = max(columns.values()) + 1

        original_rows = ws.iter_rows(
            min_row=header_row + 1,
            max_col=last_column,
        )
        cached_rows = cached_ws.iter_rows(
            min_row=header_row + 1,
            max_col=last_column,
        )

        for row_number, (original, saved) in enumerate(
            zip(original_rows, cached_rows),
            start=header_row + 1,
        ):
            # 必須3列がすべて空の行だけ無視する。
            if all(
                text(original[columns[name]].value) == ""
                for name in REQUIRED_COLUMNS
            ):
                continue

            try:
                data = {}

                for name, index in columns.items():
                    cell = original[index]
                    value = cell.value

                    if cell.data_type == "f":
                        if not allow_cached_formulas:
                            raise ValueError(
                                f"「{name}」に数式があります。"
                                "値に置換するか、"
                                "--allow-cached-formulasを指定してください"
                            )

                        value = saved[index].value
                        formula_count += 1

                        if value is None:
                            raise ValueError(
                                f"「{name}」の保存済み計算結果がありません"
                            )

                    data[name] = value

                receipt_id = parse_receipt_id(data["通し番号"])
                amount = parse_number(data["金額"], currency=True)
                team = text(data["班名"])

                if not team:
                    raise ValueError("班名が空欄です")

                groups[receipt_id].append({
                    "row": row_number,
                    "amount": amount,
                    "team": team,
                })

            except ValueError as exc:
                errors.append(f"Excel {row_number}行目: {exc}")

        fail_if_errors(errors)

        if not groups:
            raise ValueError("処理対象の行がありません")

        selected = {}
        errors = []

        for receipt_id, records in groups.items():
            maximum = max(record["amount"] for record in records)

            winners = [
                record
                for record in records
                if record["amount"] == maximum
            ]

            teams = {record["team"] for record in winners}

            if len(teams) != 1:
                details = [
                    (record["row"], record["team"])
                    for record in winners
                ]
                errors.append(
                    f"番号{receipt_id}: 最大金額{maximum}で"
                    f"班名が競合しています: {details}"
                )
                continue

            selected[receipt_id] = {
                **winners[0],
                "tied_rows": ";".join(
                    str(record["row"]) for record in winners
                ),
                # 班の表示順は、採用行ではなく元データの初出順。
                "first_row": min(record["row"] for record in records),
            }

        fail_if_errors(errors)

        if formula_count:
            print(
                "注意: 数式の保存済み計算結果を使用しました。"
                "計算結果の鮮度は保証できません。",
                file=sys.stderr,
            )

        return selected

    finally:
        formulas.close()
        if cached is not None:
            cached.close()


def match_images(folder, selected):
    """
    ファイル名が数字だけの画像を対応付ける。
    例：19.jpg、019.png

    サブフォルダも検索する。
    """
    if not folder.is_dir():
        raise ValueError(f"画像フォルダがありません: {folder}")

    images = {}
    errors = []

    for path in sorted(folder.rglob("*")):
        if not path.is_file():
            continue

        if path.suffix.lower() not in IMAGE_EXTENSIONS:
            continue

        if not re.fullmatch(r"[0-9]+", path.stem):
            errors.append(
                f"画像名は番号だけにしてください: {path}"
            )
            continue

        receipt_id = int(path.stem)

        if receipt_id <= 0:
            errors.append(f"画像番号は正の整数が必要です: {path}")
            continue

        if receipt_id in images:
            errors.append(
                f"番号{receipt_id}の画像が重複しています:\n"
                f"  {images[receipt_id]}\n"
                f"  {path}"
            )
        else:
            images[receipt_id] = path

    for receipt_id in sorted(set(selected) - set(images)):
        errors.append(
            f"番号{receipt_id}: 対応画像がありません"
        )

    for receipt_id in sorted(set(images) - set(selected)):
        errors.append(
            f"番号{receipt_id}: Excelに対応する行がありません "
            f"({images[receipt_id]})"
        )

    fail_if_errors(errors)
    return images


def image_stream(path):
    """
    EXIF回転を反映し、透過部分を白背景にしてPNG化する。
    余白の切り取りはしない。
    """
    with Image.open(path) as source:
        if getattr(source, "n_frames", 1) != 1:
            raise ValueError(
                f"複数フレーム画像は未対応です: {path}"
            )

        source.load()
        image = ImageOps.exif_transpose(source).convert("RGBA")

        background = Image.new("RGBA", image.size, "white")
        background.alpha_composite(image)

        stream = io.BytesIO()
        background.convert("RGB").save(stream, format="PNG")
        stream.seek(0)

        return stream, image.size


def add_text(slide, value, x, y, width, height, size):
    box = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )

    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.text = value

    for paragraph in frame.paragraphs:
        paragraph.font.name = "Yu Gothic"
        paragraph.font.size = Pt(size)


def build_layout(selected, dimensions, vertical_count):
    """
    縦方向の段数から画像の基準高さを決める。
    横方向は実際の画像幅で詰め、入らなければ次の段へ。
    班内の通し番号順は維持する。
    """
    available_width = SLIDE_WIDTH - 2 * MARGIN
    available_height = SLIDE_HEIGHT - CONTENT_TOP - MARGIN

    row_height = (
        available_height - VERTICAL_GAP * (vertical_count - 1)
    ) / vertical_count

    target_image_height = row_height - CAPTION_HEIGHT

    teams = defaultdict(list)

    for receipt_id, item in sorted(
        selected.items(),
        key=lambda pair: pair[1]["first_row"],
    ):
        teams[item["team"]].append(receipt_id)

    plans = []

    for team, receipt_ids in teams.items():
        team_pages = []
        current = []
        row_index = 0
        used_width = 0.0

        for receipt_id in sorted(receipt_ids):
            pixel_width, pixel_height = dimensions[receipt_id]

            scale = min(
                target_image_height / pixel_height,
                available_width / pixel_width,
            )

            width = pixel_width * scale
            height = pixel_height * scale

            gap = HORIZONTAL_GAP if used_width > 0 else 0.0

            if used_width + gap + width > available_width + 1e-9:
                row_index += 1
                used_width = 0.0
                gap = 0.0

            if row_index >= vertical_count:
                team_pages.append(current)
                current = []
                row_index = 0
                used_width = 0.0
                gap = 0.0

            x = MARGIN + used_width + gap
            y = CONTENT_TOP + row_index * (
                row_height + VERTICAL_GAP
            )

            current.append((
                receipt_id, x, y, width, height
            ))

            used_width += gap + width

        if current:
            team_pages.append(current)

        for local_page, placements in enumerate(
            team_pages, start=1
        ):
            plans.append({
                "team": team,
                "local_page": local_page,
                "page_count": len(team_pages),
                "placements": placements,
            })

    return plans


def generate(selected, images, output, vertical_count):
    csv_path = output.with_suffix(".csv")

    if output.suffix.lower() != ".pptx":
        raise ValueError("出力名の拡張子は .pptx にしてください")

    if output.exists() or csv_path.exists():
        raise ValueError(
            "同名の出力ファイルがあります。"
            "別の出力名を指定してください"
        )

    if not output.parent.is_dir():
        raise ValueError(
            f"出力先フォルダがありません: {output.parent}"
        )

    # 全画像を事前検査。
    dimensions = {}
    errors = []

    for receipt_id, path in images.items():
        try:
            stream, size = image_stream(path)
            stream.close()
            dimensions[receipt_id] = size
        except Exception as exc:
            errors.append(
                f"番号{receipt_id}: 画像を読み込めません: {exc}"
            )

    fail_if_errors(errors)

    plans = build_layout(
        selected, dimensions, vertical_count
    )

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH)
    prs.slide_height = Inches(SLIDE_HEIGHT)

    audit = []

    for plan in plans:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        page = len(prs.slides)

        add_text(
            slide,
            (
                f"{plan['team']}  "
                f"({plan['local_page']}/{plan['page_count']})"
            ),
            MARGIN,
            0.15,
            SLIDE_WIDTH - 2 * MARGIN,
            0.65,
            22,
        )

        for receipt_id, x, y, width, height in plan["placements"]:
            stream, _ = image_stream(images[receipt_id])

            try:
                slide.shapes.add_picture(
                    stream,
                    Inches(x),
                    Inches(y),
                    width=Inches(width),
                    height=Inches(height),
                )
            finally:
                stream.close()

            add_text(
                slide,
                str(receipt_id),
                x,
                y + height,
                width,
                CAPTION_HEIGHT,
                9,
            )

            item = selected[receipt_id]

            audit.append([
                receipt_id,
                item["row"],
                item["tied_rows"],
                str(item["amount"]),
                item["team"],
                str(images[receipt_id]),
                page,
            ])

    # 番号ごとに一度だけ配置されているか照合。
    placed = [record[0] for record in audit]

    if (
        len(placed) != len(selected)
        or len(set(placed)) != len(placed)
        or set(placed) != set(selected)
    ):
        raise RuntimeError("配置数の照合に失敗しました")

    # スライド外にはみ出していないか確認。
    tolerance = Inches(0.001)

    for page, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if (
                shape.left < -tolerance
                or shape.top < -tolerance
                or shape.left + shape.width
                > prs.slide_width + tolerance
                or shape.top + shape.height
                > prs.slide_height + tolerance
            ):
                raise RuntimeError(
                    f"{page}ページ目に範囲外の配置があります"
                )

    # 一時ファイルで生成し、再読込後に出力する。
    with tempfile.TemporaryDirectory(
        dir=output.parent
    ) as temporary:
        temporary = Path(temporary)
        pptx_tmp = temporary / "result.pptx"
        csv_tmp = temporary / "result.csv"

        prs.save(pptx_tmp)

        with csv_tmp.open(
            "w", encoding="utf-8-sig", newline=""
        ) as file:
            writer = csv.writer(file)
            writer.writerow([
                "通し番号",
                "採用Excel行",
                "最大金額同額行",
                "最大金額",
                "班名",
                "画像ファイル",
                "配置ページ",
            ])
            writer.writerows(audit)

        reopened = Presentation(pptx_tmp)

        picture_count = sum(
            shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            for slide in reopened.slides
            for shape in slide.shapes
        )

        if picture_count != len(selected):
            raise RuntimeError("保存後の画像数が一致しません")

        # 既存ファイルは上書きしない。
        # 2ファイルの出力に失敗した場合は、今回作成分を削除する。
        created = []
        try:
            for source, destination in (
                (pptx_tmp, output),
                (csv_tmp, csv_path),
            ):
                with destination.open("xb") as target:
                    created.append(destination)
                    with source.open("rb") as original:
                        import shutil
                        shutil.copyfileobj(original, target)
        except BaseException:
            for destination in created:
                destination.unlink(missing_ok=True)
            raise

    print(
        f"完了: レシート{len(selected)}枚 / "
        f"{len(prs.slides)}ページ"
    )
    print(f"PowerPoint: {output}")
    print(f"照合CSV: {csv_path}")


def main():
    parser = argparse.ArgumentParser(
        description=(
            "通し番号ごとの最大金額の行の班名で"
            "レシート画像を分類し、PowerPointを生成する"
        )
    )

    parser.add_argument(
        "--excel", type=Path, required=True
    )
    parser.add_argument(
        "--images", type=Path, required=True
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("receipts.pptx"),
    )
    parser.add_argument(
        "--sheet",
        help="対象シート名。シートが1つの場合は省略可能",
    )
    parser.add_argument(
        "--header-row",
        type=int,
        default=1,
        help="列名が書かれている行番号。既定: 1",
    )
    parser.add_argument(
        "--rows",
        type=int,
        choices=range(1, 7),
        default=3,
        help="縦方向の段数。既定: 3",
    )
    parser.add_argument(
        "--allow-cached-formulas",
        action="store_true",
        help="数式セルの保存済み計算結果を使用する",
    )

    args = parser.parse_args()

    if args.header_row < 1:
        parser.error("--header-rowは1以上にしてください")

    try:
        selected = read_excel(
            args.excel,
            args.sheet,
            args.header_row,
            args.allow_cached_formulas,
        )

        images = match_images(
            args.images, selected
        )

        generate(
            selected,
            images,
            args.output,
            args.rows,
        )

    except Exception as exc:
        print(f"処理失敗:\n{exc}", file=sys.stderr)
        return 1

    return 0


if __name__ == "__main__":
    sys.exit(main())